"""Export the BMW-provided workbooks into research contract v1 batches.

One batch per artifact or sheet: customer feedback per vehicle (one record per feedback ID, every
classification row kept as an original label), the CN/EU and US customer-study sheets (one metric
per attribute and cohort) and sales volumes (one metric per vehicle sheet, market and year).

Originals are read from the ignored data/ directory and stay unchanged. Output goes to the ignored
outputs/research/submissions/bmw/ unless --output names another directory. Missing BMW VFC topics
are appended to research/catalog/features.json so feedback labels map to feature keys. Nothing is
sent to a model; permissions and review status stay unresolved unless explicitly set.
"""

import argparse
from collections import OrderedDict, defaultdict
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
FEEDBACK_MODELS = ("F70", "G60", "G70")
FEEDBACK_SHEET = "Feedback_Explorer"
FEEDBACK_CONTEXT = ("Source", "Brand", "Sales Description", "Country", "Derivate (e-code)", "Engine Type",
                    "Customer Feedback")
# Per classification row, in this order: the backend pairs "Feedback Type" and "Vfc level3 Name"
# with the most recent "Vfc level2 Name", so the topic must open each group.
FEEDBACK_LABELS = ("Vfc level2 Name", "Feedback Type", "Vfc level3 Name")
STUDY_FILE = "F70_G60_G68_G70_customer_studies.xlsx"
SALES_FILE = "sales_volumes.xlsx"
US_SCALE = ("I Hate It", "A Failure", "Unsatisfactory", "Satisfactory", "Excellent", "Delightful", "I Love It")
STUDY_MARKETS = {"China": "CN", "EU": "EU", "USA": "US", "US": "US"}
SALES_SUBJECTS = {"F70": "bmw:F70", "G60_G68": "bmw:G60_G68", "G70": "bmw:G70"}
SALES_FORECAST_YEARS = {2030}
NOT_A_FEATURE = re.compile(r"^(no_class_found|general enquiry\b)", re.IGNORECASE)
MARKET_CODE = re.compile(r"^[A-Z]{2}$")
UNRESOLVED_PERMISSION = ("Unresolved: BMW-provided challenge data. Keep out of Git and external LLMs "
                         "until the team decides.")

# Appended VFC topics get an area from the existing catalog vocabulary: explicit entries first,
# then the first matching keyword rule. Unmatched topics keep a null area.
AREA_OVERRIDES = {
    "12 V connection": "space_practicality", "Bumper, front": "exterior_design", "Bumper, rear": "exterior_design",
    "Hood": "exterior_design", "Hood, open / close": "access_locking", "Running board": "exterior_design",
    "Towing eye": "exterior_design", "Windshield": "lighting_visibility", "Horns": "safety",
    "Side View": "adas_parking", "Night Vision": "adas_parking", "Map update": "hmi_infotainment",
    "Smart POI search": "hmi_infotainment", "Touch Command": "hmi_infotainment",
    "Video: playback / streaming": "hmi_infotainment", "My BMW / MINI Portal": "connectivity_app_digital",
    "Teleservices": "connectivity_app_digital", "Sun protection": "interior_design_quality",
    "Sun visor, front": "interior_design_quality", "Caring Car": "seats_comfort",
    "Interior mirror, garage door opener": "connectivity_app_digital",
    "Vehicle interior, lighting": "interior_design_quality", "Heated rear window": "climate",
    "Trunk light": "lighting_visibility", "Start engine/switch on drive readiness": "powertrain_performance",
    "Display key": "access_locking",
}
AREA_RULES = (
    ("range_charging_efficiency", r"range|charg|battery|recuperation|energy|consumption|efficien"),
    ("adas_parking", r"assist|park|camera|cruise|lane|collision|warning|recorder|traffic|all-round view|"
                     r"back-up|distance|autonomous|automated"),
    ("connectivity_app_digital", r"\bapp\b|apps|digital key|remote|esim|wi-?fi|connected|software|upgrade|rsu|"
                                 r"android|carplay|bluetooth|smartphone|phone|usb|online|account"),
    ("hmi_infotainment", r"display|screen|navigation|voice|gesture|idrive|controller|profile|operating|menu|"
                         r"infotainment|radio|media|head-up|mode"),
    ("seats_comfort", r"seat|lumbar|massage|headrest|steering wheel heat"),
    ("access_locking", r"door|trunk|lid|tailgate|\bkey\b|lock|alarm|flap|window|opening|closing|start engine"),
    ("lighting_visibility", r"light|lamp|mirror|wiper|washer|visibility"),
    ("climate", r"climate|air|heat|vent|blower|cooling|temperature"),
    ("audio_acoustics", r"sound|audio|speaker|noise|acoustic|loudspeaker"),
    ("powertrain_performance", r"engine|drive|gear|shift|paddle|launch|pedal|transmission|exhaust|fuel"),
    ("space_practicality", r"glove|storage|cupholder|mat|console|armrest|space|compartment|holder|luggage"),
    ("interior_design_quality", r"interior|trim|ambient|roof|upholster|leather|material"),
    ("exterior_design", r"exterior|paint|wheel|rim|design"),
    ("ride_handling", r"tire|tyre|tirefit|suspension|steering|brak|chassis|handling"),
    ("safety", r"emergency|airbag|safety|belt|sos|crash"),
)


def slug(text):
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def norm_label(text):
    # Same normalization the backend uses to join VFC topics to catalog labels.
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def snapshot_id(record):
    """Readable observation key plus a content hash: identical reruns keep the ID, changed content
    or metadata creates a new snapshot. Row positions are excluded so a re-sorted file keeps IDs."""
    content = {key: record[key] for key in record if key not in ("record_id", "source_id", "locator")}
    digest = hashlib.sha256(json.dumps(content, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return f"{record['dedup_key']}:{digest[:12]}"


def new_record(dedup_key, kind, locator, subject_ids, markets, *, text=None, language=None, metric=None,
               source_fields=None):
    record = {
        "document_type": "record", "schema_version": "1.0", "record_id": None, "source_id": None,
        "dedup_key": dedup_key, "record_kind": kind, "locator": locator,
        "content_date": None, "content_datetime": None,
        "subject_ids": subject_ids, "markets": markets,
        "text_original": text, "language": language, "translation_en": None,
        "metric": metric, "availability": None, "source_fields": source_fields or {},
    }
    record["record_id"] = snapshot_id(record)
    return record


def metric(name, value, unit, statistic, period, population):
    return {"name": name, "value": value, "unit": unit, "statistic": statistic, "period": period,
            "population": population, "sample_base": None, "scale": None}


def sheet_rows(path, sheet, data_only=True):
    workbook = load_workbook(path, read_only=True, data_only=data_only)
    try:
        return [tuple(row) for row in workbook[sheet].iter_rows(values_only=True)]
    finally:
        workbook.close()


class Batch:
    def __init__(self, topic, dataset_key, channel, title, artifact, records, *, retrieved_count, method,
                 limitations, notes):
        self.topic = topic
        self.dataset_key = dataset_key
        self.channel = channel
        self.title = title
        self.artifact = artifact
        self.records = records
        self.retrieved_count = retrieved_count
        self.method = method
        self.limitations = limitations
        self.notes = notes

    def source(self, options, source_id, collected_at):
        return {
            "document_type": "source", "schema_version": "1.0",
            "source_id": source_id, "dataset_key": self.dataset_key,
            "stream": "bmw", "channel": self.channel, "title": self.title,
            "locator": f"data/{self.artifact.name}", "collected_at": collected_at,
            "collector": options.collector, "is_example": False,
            "normalized_record_count": len(self.records),
            "sampling": {
                "method": self.method, "query": None, "date_start": None, "date_end": None,
                "retrieved_count": self.retrieved_count, "complete": True, "limitations": self.limitations,
            },
            "provenance": {
                "origin": "bmw_provided", "verification": "provenance_checked",
                "notes": ("BMW challenge Google Drive folder (docs/datasets/README.md); file hash recorded, "
                          "content not independently verified. Exported by scripts/export_bmw.py."),
            },
            "artifact": {"path": f"data/{self.artifact.name}", "sha256": sha256_file(self.artifact)},
            "export_permissions": {
                "git": options.git, "external_llm": options.external_llm,
                "notes": options.permission_note or UNRESOLVED_PERMISSION,
            },
            "review_status": options.review_status,
        }


# ---------------------------------------------------------------- feedback


def export_feedback(data_dir, model):
    path = data_dir / f"{model}_feedback_hackathon.xlsx"
    rows = sheet_rows(path, FEEDBACK_SHEET)
    header = [clean(cell) for cell in rows[0]]
    index = {name: position for position, name in enumerate(header) if name}
    missing = [name for name in ("ID",) + FEEDBACK_CONTEXT + FEEDBACK_LABELS if name not in index]
    if missing:
        raise ValueError(f"{path}: missing columns {missing}")
    columns = sorted({get_column_letter(index[name] + 1) for name in ("ID",) + FEEDBACK_CONTEXT + FEEDBACK_LABELS},
                     key=lambda letter: (len(letter), letter))

    groups = OrderedDict()
    for row_number, row in enumerate(rows[1:], 2):
        feedback_id = clean(row[index["ID"]])
        if feedback_id is None:
            if any(clean(cell) for cell in row):
                raise ValueError(f"{path}:{row_number}: classification row without an ID")
            continue
        groups.setdefault(feedback_id, []).append((row_number, row))

    records, conflicts, issues = [], [], defaultdict(list)
    for feedback_id, members in groups.items():
        contexts = {tuple(row[index[name]] for name in FEEDBACK_CONTEXT) for _, row in members}
        if len(contexts) > 1:
            conflicts.append({"id": feedback_id, "rows": [number for number, _ in members]})
            continue
        first = members[0][1]
        text = first[index["Customer Feedback"]]
        if not isinstance(text, str) or not text.strip():
            issues["empty_text"].append(feedback_id)
            continue
        derivative = clean(first[index["Derivate (e-code)"]])
        if derivative != model:
            issues["derivative_mismatch"].append(feedback_id)
        country = clean(first[index["Country"]])
        markets = [country] if country and MARKET_CODE.match(country) else []
        if not markets:
            issues["unresolved_country"].append(feedback_id)
        labels = [{"name": "Engine Type", "value": clean(first[index["Engine Type"]])}]
        for _, row in members:
            labels += [{"name": name, "value": clean(row[index[name]])} for name in FEEDBACK_LABELS]
        records.append(new_record(
            f"bmw-feedback:{model}:{feedback_id}", "feedback",
            {"kind": "spreadsheet", "sheet": FEEDBACK_SHEET, "rows": [number for number, _ in members],
             "columns": columns},
            [f"bmw:{derivative or model}"], markets, text=text, language="en",
            source_fields={"Source": clean(first[index["Source"]]), "original_labels": labels},
        ))

    classification_rows = sum(len(members) for members in groups.values())
    topics = defaultdict(int)
    for _, members in groups.items():
        for _, row in members:
            topic = clean(row[index["Vfc level2 Name"]])
            if topic:
                topics[topic] += 1
    by_source, by_country, by_engine = defaultdict(int), defaultdict(int), defaultdict(int)
    for record in records:
        by_source[record["source_fields"]["Source"]] += 1
        by_country[(record["markets"] or ["unresolved"])[0]] += 1
        by_engine[record["source_fields"]["original_labels"][0]["value"]] += 1
    limitations = [
        "One spreadsheet row per BMW VFC classification; rows were collapsed to one record per feedback ID "
        "after checking that ID, text and context agree. All contributing rows and labels are preserved.",
        "Feedback sources A-D are anonymized by BMW; their channels are not documented.",
        "No dates in the workbook; content_date stays null.",
        "No China and no Germany feedback; country coverage is unbalanced and is not a market sample.",
        "Texts are English; whether BMW translated non-English originals is not documented.",
        "VFC labels are BMW's classification of unknown method; derived claims from them stay pending.",
    ]
    if conflicts:
        limitations.append(f"{len(conflicts)} feedback IDs with disagreeing rows were excluded; see notes.md.")
    return Batch(
        f"bmw-feedback-{model.lower()}", f"bmw-feedback:{model}", "bmw_feedback",
        f"BMW customer feedback, {model} (Feedback_Explorer export)", path, records,
        retrieved_count=classification_rows,
        method="Full export of the BMW-provided feedback workbook; one record per feedback ID.",
        limitations=limitations,
        notes={
            "classification_rows": classification_rows, "unique_ids": len(groups), "records": len(records),
            "conflicting_ids": conflicts, "issues": {key: value for key, value in issues.items()},
            "by_source": dict(sorted(by_source.items())), "by_country": dict(sorted(by_country.items(), key=lambda kv: -kv[1])),
            "by_engine": dict(sorted(by_engine.items())), "vfc_topics": dict(topics), "columns": columns,
        },
    )


# ---------------------------------------------------------------- customer study


def vehicle_subject(label):
    match = re.search(r"\b([FG]\d{2})\b", label or "")
    if not match:
        raise ValueError(f"unrecognized study vehicle label {label!r}")
    return f"bmw:{match.group(1)}"


def export_study_cn_eu(data_dir, sheet="CN_EU_2025"):
    path = data_dir / STUDY_FILE
    rows = sheet_rows(path, sheet)
    country_row = next(position for position, row in enumerate(rows) if any(clean(c) in STUDY_MARKETS for c in row[1:]))
    vehicle_row = country_row + 1
    # The country header is a merged cell: forward-fill it across the vehicle columns.
    cohorts, current = {}, None
    for column, cell in enumerate(rows[vehicle_row][1:], 1):
        country = clean(rows[country_row][column]) if column < len(rows[country_row]) else None
        current = country or current
        vehicle = clean(cell)
        if vehicle:
            if current not in STUDY_MARKETS:
                raise ValueError(f"{path}:{sheet}: vehicle column {column + 1} has no market header")
            cohorts[column] = (STUDY_MARKETS[current], current, vehicle, vehicle_subject(vehicle))

    records, attribute, attribute_row, means = [], None, None, 0
    for row_number, row in enumerate(rows[vehicle_row:], vehicle_row + 1):
        label = clean(row[0])
        if label is None:
            continue
        if label.lower() != "mean":
            # The first attribute shares the vehicle header row.
            attribute, attribute_row = label, row_number
            continue
        if attribute is None:
            raise ValueError(f"{path}:{sheet}:{row_number}: Mean row without an attribute")
        means += 1
        for column, (market, country, vehicle, subject) in cohorts.items():
            value = row[column] if column < len(row) else None
            if value is None or str(value).strip() == "":
                continue
            records.append(new_record(
                f"bmw-study:{sheet}:{market}:{subject.split(':')[1]}:{slug(attribute)}", "metric",
                {"kind": "spreadsheet", "sheet": sheet, "rows": sorted({attribute_row, row_number}),
                 "columns": [get_column_letter(column + 1)]},
                [subject], [market],
                metric=metric(attribute, float(value), "mean score", "mean", "2025",
                              f"BMW customer study 2025, {country}, owners of {vehicle}; respondent definition "
                              "and sample size not given"),
                source_fields={"original_labels": [{"name": "Dataset country", "value": country},
                                                   {"name": "Vehicle", "value": vehicle}]},
            ))
    values = [r["metric"]["value"] for r in records]
    return Batch(
        "bmw-study-cn-eu-2025", f"bmw-study:{sheet}", "bmw_customer_study",
        "BMW customer study 2025, China and EU (mean scores)", path, records,
        retrieved_count=means * len(cohorts),
        method=f"Full export of sheet {sheet}; one metric per attribute and market/vehicle column.",
        limitations=[
            "Mean scores; the scale is not stated. Values between "
            f"{min(values):.2f} and {max(values):.2f} suggest a 10-point scale (inferred, scale left null).",
            "No sample sizes; differences are relative signals, not tested ones.",
            "Not comparable with the US_2025 distributions (different scale and question set).",
            "China covers the 5 Series long wheelbase G68 only; G60 appears only in the EU column.",
        ],
        notes={"attributes": means, "cohorts": [list(c) for c in cohorts.values()], "records": len(records)},
    )


def export_study_us(data_dir, sheet="US_2025"):
    path = data_dir / STUDY_FILE
    rows = sheet_rows(path, sheet)
    vehicles = {column: (clean(cell), vehicle_subject(clean(cell))) for column, cell in enumerate(rows[0])
                if column >= 2 and clean(cell)}
    blocks, attribute, issues = [], None, defaultdict(list)
    for row_number, row in enumerate(rows, 1):
        label, base = clean(row[0]), clean(row[1]) if len(row) > 1 else None
        if label is None:
            continue
        if base is None:
            attribute = {"name": label, "row": row_number, "values": {}, "sample_total": None}
            blocks.append(attribute)
        elif attribute is None:
            raise ValueError(f"{path}:{sheet}:{row_number}: value row before the first attribute")
        elif label == "Sample total":
            attribute["sample_total"] = {column: row[column] for column in vehicles}
        elif label in US_SCALE:
            attribute["values"][label] = (row_number, {column: row[column] for column in vehicles})
        else:
            issues["unknown_row_label"].append(row_number)

    names = defaultdict(int)
    for block in blocks:
        names[block["name"]] += 1
    records = []
    for block in blocks:
        key_name = block["name"] if names[block["name"]] == 1 else f"{block['name']} (row {block['row']})"
        if set(block["values"]) != set(US_SCALE):
            issues["incomplete_block"].append(block["row"])
        for column, (vehicle, subject) in vehicles.items():
            total = sum(float(values[column]) for _, values in block["values"].values() if values[column] is not None)
            if abs(total - 1.0) > 0.01:
                issues["block_sum_not_1"].append({"row": block["row"], "column": get_column_letter(column + 1),
                                                  "sum": round(total, 4)})
            for label in US_SCALE:
                if label not in block["values"]:
                    continue
                row_number, values = block["values"][label]
                if values[column] is None:
                    continue
                records.append(new_record(
                    f"bmw-study:{sheet}:US:{subject.split(':')[1]}:{slug(key_name)}:{slug(label)}", "metric",
                    {"kind": "spreadsheet", "sheet": sheet, "rows": sorted({block["row"], row_number}),
                     "columns": [get_column_letter(column + 1)]},
                    [subject], ["US"],
                    # "Attribute | Scale label" is the name format the backend reads into a weighted mean.
                    metric=metric(f"{key_name} | {label}", float(values[column]), "share of respondents",
                                  "proportion", "2025",
                                  f"BMW customer study 2025, USA, owners of {vehicle}; sample size not given"),
                    source_fields={"original_labels": [
                        {"name": "Vehicle", "value": vehicle}, {"name": "Response", "value": label},
                        {"name": "Response scale", "value": " < ".join(US_SCALE)},
                        {"name": "Base", "value": clean(rows[row_number - 1][1])},
                    ]},
                ))
    return Batch(
        "bmw-study-us-2025", f"bmw-study:{sheet}", "bmw_customer_study",
        "BMW customer study 2025, USA (7-point distributions)", path, records,
        retrieved_count=len(blocks) * len(US_SCALE) * len(vehicles),
        method=f"Full export of sheet {sheet}; one proportion per attribute, response level and vehicle.",
        limitations=[
            "Share of respondents per level of a 7-point scale (I Hate It ... I Love It); no sample sizes.",
            "The 'Sample total' row is always 1 (100 %), not a respondent count; it was not exported.",
            "Not comparable with the CN/EU mean scores (different scale and question set).",
            "No US study for the 1 Series F70.",
        ],
        notes={"attributes": len(blocks), "vehicles": [v for v, _ in vehicles.values()], "records": len(records),
               "duplicate_attribute_names": [n for n, c in names.items() if c > 1],
               "issues": {key: value for key, value in issues.items()}},
    )


# ---------------------------------------------------------------- sales volumes


def export_sales(data_dir):
    path = data_dir / SALES_FILE
    workbook = load_workbook(path, read_only=True)
    sheets = workbook.sheetnames
    workbook.close()
    records, formulas, unknown, cells = [], [], [], 0
    for sheet in sheets:
        if sheet not in SALES_SUBJECTS:
            unknown.append(sheet)
            continue
        values = sheet_rows(path, sheet, data_only=True)
        raw = sheet_rows(path, sheet, data_only=False)
        header = [clean(cell) for cell in values[0]]
        year_columns = {column: int(name.rsplit("_", 1)[1]) for column, name in enumerate(header)
                        if name and re.fullmatch(r"volume_\d{4}", name)}
        for row_number, row in enumerate(values[1:], 2):
            market, code = clean(row[0]), clean(row[1])
            if not market:
                continue
            for column, year in year_columns.items():
                cells += 1
                value = row[column]
                formula = raw[row_number - 1][column]
                labels = [{"name": "market", "value": market}, {"name": "market_code", "value": code}]
                if isinstance(formula, str) and formula.startswith("="):
                    formulas.append({"sheet": sheet, "cell": f"{get_column_letter(column + 1)}{row_number}",
                                     "formula": formula, "cached_value": value})
                    labels.append({"name": "cell formula", "value": formula})
                if value is None:
                    continue
                forecast = year in SALES_FORECAST_YEARS
                records.append(new_record(
                    f"bmw-sales:{sheet}:{code}:{year}", "metric",
                    {"kind": "spreadsheet", "sheet": sheet, "rows": [row_number],
                     "columns": [get_column_letter(column + 1)]},
                    [SALES_SUBJECTS[sheet]], [code] if code and MARKET_CODE.match(code) else [],
                    metric=metric(f"Sales volume {year}", float(value), "vehicles",
                                  "forecast" if forecast else "count", str(year),
                                  f"{market} ({code}), sheet {sheet}; all powertrains combined"),
                    source_fields={"original_labels": labels},
                ))
    limitations = [
        "No powertrain split, although the BMW data slide describes one.",
        "5 Series G60 and G68 are combined in sheet G60_G68 and assigned to bmw:G60_G68, never to G60 alone.",
        "Whether 2024/2025 are actuals or plan values is not documented; 2030 lies in the future and is "
        "exported as a forecast.",
        "Rest of World (RoW) is not a country or EU code; those records keep an empty markets list.",
    ]
    if formulas:
        limitations.append(f"{len(formulas)} cell(s) hold formulas; their cached values were exported (see notes.md).")
    return Batch(
        "bmw-sales-volumes", "bmw-sales:volumes", "bmw_sales",
        "BMW sales volumes 2024, 2025 and 2030 by market", path, records,
        retrieved_count=cells, method="Full export of all vehicle sheets; one metric per market and year.",
        limitations=limitations,
        notes={"sheets": sheets, "records": len(records), "formulas": formulas, "unknown_sheets": unknown},
    )


# ---------------------------------------------------------------- catalog


def area_for(topic):
    if topic in AREA_OVERRIDES:
        return AREA_OVERRIDES[topic]
    text = topic.lower()
    for area, pattern in AREA_RULES:
        if re.search(pattern, text):
            return area
    return None


def extend_catalog(catalog_path, topic_counts):
    """Append VFC level2 topics that no catalog label matches. Existing entries are never changed."""
    document = json.loads(catalog_path.read_text(encoding="utf-8"))
    features = document["features"]
    labels = {norm_label(entry.get("label", "")) for entry in features}
    keys = {entry["feature_key"] for entry in features}
    added, skipped = [], []
    for topic in sorted(topic_counts, key=lambda t: (-topic_counts[t], t)):
        if NOT_A_FEATURE.match(topic):
            skipped.append(topic)
            continue
        if norm_label(topic) in labels:
            continue
        key = slug(topic)
        if key in keys:
            key = f"{key}_vfc"
        entry = {"feature_key": key, "label": topic, "area": area_for(topic),
                 "origin": "BMW VFC level2 topic (appended by scripts/export_bmw.py; area by keyword rule)"}
        features.append(entry)
        labels.add(norm_label(topic))
        keys.add(key)
        added.append(entry)
    if added:
        catalog_path.write_text(catalog_text(document), encoding="utf-8")
    return added, skipped


def catalog_text(document):
    # The catalog's own layout: top-level keys on their own lines, one feature object per line.
    head = "".join(f"  {json.dumps(key)}: {json.dumps(value, ensure_ascii=False)},\n"
                   for key, value in document.items() if key != "features")
    body = ",\n".join("    " + json.dumps(entry, ensure_ascii=False) for entry in document["features"])
    return "{\n" + head + '  "features": [\n' + body + "\n  ]\n}\n"


# ---------------------------------------------------------------- output


def write_batch(output, batch, options, collected_at):
    source_id = f"{options.collector}-{options.date}-{batch.topic}"
    for stale in output.glob(f"{options.collector}-*-{batch.topic}"):
        if stale.is_dir() and stale.name != source_id:
            shutil.rmtree(stale)
    directory = output / source_id
    directory.mkdir(parents=True, exist_ok=True)
    for record in batch.records:
        record["source_id"] = source_id
    source = batch.source(options, source_id, collected_at)
    (directory / "source.json").write_text(json.dumps(source, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    with (directory / "records.jsonl").open("w", encoding="utf-8") as stream:
        for record in batch.records:
            stream.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")
    (directory / "notes.md").write_text(batch_notes(batch, source), encoding="utf-8")
    return directory


def batch_notes(batch, source):
    lines = [f"# {source['title']}", "",
             f"Exported by `scripts/export_bmw.py` from `{source['artifact']['path']}` "
             f"(sha256 `{source['artifact']['sha256'][:16]}…`). Dataset key `{batch.dataset_key}`.", "",
             "## Counts", ""]
    for key, value in batch.notes.items():
        if isinstance(value, (int, float, str)):
            lines.append(f"- {key.replace('_', ' ')}: {value}")
    for key in ("by_source", "by_country", "by_engine"):
        if key in batch.notes:
            lines.append(f"- {key.replace('_', ' ')}: " + ", ".join(f"{k} {v}" for k, v in batch.notes[key].items()))
    lines += ["", "## Limitations", ""] + [f"- {item}" for item in batch.limitations]
    unresolved = []
    if batch.notes.get("conflicting_ids"):
        unresolved.append("Feedback IDs with disagreeing rows (excluded): "
                          + ", ".join(c["id"] for c in batch.notes["conflicting_ids"]))
    for key, value in (batch.notes.get("issues") or {}).items():
        if value:
            unresolved.append(f"{key.replace('_', ' ')}: {json.dumps(value)[:400]}")
    for formula in batch.notes.get("formulas") or []:
        unresolved.append(f"Formula in {formula['sheet']}!{formula['cell']} ({formula['formula']}); "
                          f"cached value {formula['cached_value']} exported.")
    if batch.notes.get("duplicate_attribute_names"):
        unresolved.append("Repeated attribute names (row number added to the key): "
                          + ", ".join(batch.notes["duplicate_attribute_names"]))
    lines += ["", "## Unresolved", ""] + ([f"- {item}" for item in unresolved] or ["- None found by the exporter."])
    permissions = source["export_permissions"]
    lines += ["", "## Review", "",
              f"- Review status: `{source['review_status']}`. Permissions: git `{json.dumps(permissions['git'])}`, "
              f"external LLM `{json.dumps(permissions['external_llm'])}`. {permissions['notes']}",
              "- Validate with `python scripts/validate_research.py <output root>`.", ""]
    return "\n".join(lines)


def parse_permission(value):
    return {"null": None, "true": True, "false": False}[value]


def build(options):
    data_dir, output = options.data_dir, options.output
    batches = []
    if "feedback" in options.datasets:
        batches += [export_feedback(data_dir, model) for model in options.feedback_models]
    if "study" in options.datasets:
        batches += [export_study_cn_eu(data_dir), export_study_us(data_dir)]
    if "sales" in options.datasets:
        batches.append(export_sales(data_dir))

    topics = defaultdict(int)
    for batch in batches:
        for topic, count in batch.notes.get("vfc_topics", {}).items():
            topics[topic] += count
    added, skipped = extend_catalog(options.catalog, topics) if topics else ([], [])

    output.mkdir(parents=True, exist_ok=True)
    collected_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    directories = [write_batch(output, batch, options, collected_at) for batch in batches]
    report = {
        "collected_at": collected_at,
        "batches": [{"source_id": d.name, "dataset_key": b.dataset_key, "records": len(b.records),
                     "retrieved_count": b.retrieved_count,
                     **{k: v for k, v in b.notes.items() if k in ("classification_rows", "unique_ids", "conflicting_ids")}}
                    for d, b in zip(directories, batches)],
        "catalog_added": [entry["feature_key"] for entry in added],
        "catalog_skipped_topics": skipped,
    }
    (output / "export-report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return report


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--data-dir", type=Path, default=ROOT / "data")
    parser.add_argument("--output", type=Path, default=ROOT / "outputs/research/submissions/bmw")
    parser.add_argument("--catalog", type=Path, default=ROOT / "research/catalog/features.json")
    parser.add_argument("--datasets", nargs="+", choices=("feedback", "study", "sales"),
                        default=["feedback", "study", "sales"])
    parser.add_argument("--feedback-models", nargs="+", choices=FEEDBACK_MODELS, default=["G60"])
    parser.add_argument("--collector", default="luis")
    parser.add_argument("--date", default=date.today().strftime("%Y%m%d"), help="batch date in source IDs")
    parser.add_argument("--git", type=parse_permission, choices=(None, True, False), default=None,
                        metavar="{null,true,false}")
    parser.add_argument("--external-llm", type=parse_permission, choices=(None, True, False), default=None,
                        metavar="{null,true,false}")
    parser.add_argument("--permission-note", default=None, help="who decided the permissions, and when")
    parser.add_argument("--review-status", choices=("pending", "accepted", "rejected"), default="pending")
    options = parser.parse_args(argv)
    if (options.git or options.external_llm) and not options.permission_note:
        parser.error("--permission-note is required when a permission is granted")
    return options


def main(argv=None):
    options = parse_args(argv)
    try:
        report = build(options)
    except (OSError, KeyError, ValueError) as error:
        print(f"Export failed: {error}", file=sys.stderr)
        return 1
    for batch in report["batches"]:
        extra = (f" from {batch['classification_rows']} classification rows" if "classification_rows" in batch else "")
        print(f"{batch['source_id']}: {batch['records']} records{extra}")
    print(f"Catalog: {len(report['catalog_added'])} VFC topics appended. Output: {options.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
